import asyncio
import gzip
import socket
from io import BytesIO

import httpx
import pytest
from bs4 import BeautifulSoup

import app.services as services
from app.ai import AIServiceError, SenseNovaClient
from app.config import Settings
from app.services import (
    JD_EMPTY_ERROR,
    RESUME_DESIGN_THEMES,
    _extract_gxrc_job,
    extract_document_text,
    fallback_resume_design,
    normalize_resume_content,
    resolve_resume_design,
    validate_jd_result,
)
from docx import Document
from app.storage import ObjectStore

PUBLIC_TEST_IP = "93.184.216.34"


def test_colored_resume_themes_use_one_primary_color_family() -> None:
    for theme in RESUME_DESIGN_THEMES.values():
        if theme["colored"]:
            assert theme["ribbon"] == theme["primary"]


def test_docx_text_extraction_ignores_duplicate_merged_cells() -> None:
    document = Document()
    table = document.add_table(rows=1, cols=2)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "2024.05-2026.01 Example Company Sales Manager"
    output = BytesIO()
    document.save(output)

    text = extract_document_text("resume.docx", output.getvalue())

    assert text.splitlines() == ["2024.05-2026.01 Example Company Sales Manager"]


def test_work_experience_is_sorted_by_latest_end_date() -> None:
    content = normalize_resume_content(
        {
            "experience": [
                {"company": "无日期公司", "period": "", "role": "顾问"},
                {"company": "早期公司", "period": "2018.01-2020.02", "role": "专员"},
                {"company": "当前公司", "period": "2025年3月 至今", "role": "经理"},
                {"company": "最近公司", "period": "2024.05-2026.01", "role": "主管"},
                {"company": "中期公司", "period": "2022.12-2023.07", "role": "工程师"},
            ]
        }
    )

    assert [item["company"] for item in content["experience"]] == [
        "当前公司",
        "最近公司",
        "中期公司",
        "早期公司",
        "无日期公司",
    ]


def test_resume_normalization_preserves_original_experience_periods() -> None:
    original = {
        "experience": [
            {"company": "广西示例科技有限公司", "period": "2021.03-2024.06", "role": "销售主管"},
            {"company": "南宁样例商贸有限公司", "period": "2018.07-2021.02", "role": "销售专员"},
        ],
        "projects": [{"project": "渠道增长项目", "period": "2023.01-2023.12", "role": "负责人"}],
        "education": [{"school": "广西大学", "period": "2014.09-2018.06", "role": "本科 · 市场营销"}],
    }
    candidate = {
        "experience": [
            {"company": "广西示例科技有限公司", "role": "销售主管", "details": ["负责重点客户开发"]},
            {"company": "南宁样例商贸有限公司", "role": "销售专员", "details": ["维护区域客户"]},
        ],
        "projects": [{"project": "渠道增长项目", "role": "负责人", "details": ["推动渠道增长"]}],
        "education": [{"school": "广西大学", "role": "本科 · 市场营销"}],
    }

    content = normalize_resume_content(candidate, original)

    assert [item["period"] for item in content["experience"]] == ["2021.03-2024.06", "2018.07-2021.02"]
    assert content["projects"][0]["period"] == "2023.01-2023.12"
    assert content["education"][0]["period"] == "2014.09-2018.06"


def test_resume_design_prefers_manual_theme_and_validates_ai_plan() -> None:
    jd = {"title": "高级运维工程师", "keywords": ["生产设备", "制造", "质量"]}
    manual = fallback_resume_design(jd, "ats_mono")
    planned = resolve_resume_design(
        jd,
        planned={"theme_id": "operations_terra", "layout_id": "timeline_focus", "reason": "制造现场岗位"},
    )

    assert manual["theme_id"] == "ats_mono"
    assert manual["label"] == "打印友好"
    assert planned["theme_id"] == "operations_terra"
    assert planned["layout_id"] == "timeline_focus"
    assert planned["source"] == "ai"


def _mock_public_dns(monkeypatch) -> None:
    def fake_getaddrinfo(host: str, port: int, *_args, **_kwargs):
        address = "127.0.0.1" if host == "127.0.0.1" else PUBLIC_TEST_IP
        return [(socket.AF_INET, socket.SOCK_STREAM, 6, "", (address, port))]

    monkeypatch.setattr(services.socket, "getaddrinfo", fake_getaddrinfo)


def _mock_async_transport(monkeypatch, handler) -> None:
    real_async_client = httpx.AsyncClient
    transport = httpx.MockTransport(handler)

    def client_factory(*args, **kwargs):
        return real_async_client(*args, transport=transport, **kwargs)

    monkeypatch.setattr(services.httpx, "AsyncClient", client_factory)


class OversizedStream(httpx.AsyncByteStream):
    def __init__(self) -> None:
        self.closed = False

    async def __aiter__(self):
        yield b"a" * 1_500_000
        yield b"b" * 600_000

    async def aclose(self) -> None:
        self.closed = True


class StaticStream(httpx.AsyncByteStream):
    def __init__(self, content: bytes) -> None:
        self.content = content

    async def __aiter__(self):
        yield self.content

    async def aclose(self) -> None:
        pass


def test_extract_gxrc_job_uses_server_rendered_detail() -> None:
    html = """
    <html><head><title>业务经理职位信息_示例科技有限公司 - 广西人才网</title></head>
    <body>
      <h1>业务经理</h1>
      <div class="ent-name"><a>示例科技有限公司</a></div>
      <span class="layui-keyword-text">产教融合</span>
      <span class="layui-keyword-text">校企合作</span>
      <pre id="examineSensitiveWordsContent">任职要求：
1.三年以上相关经验；
2.具备商务谈判能力。
岗位职责：
1.负责市场拓展；
2.推进项目落地。</pre>
      <footer>推荐职位和网站页脚不应进入结果</footer>
    </body></html>
    """
    result = _extract_gxrc_job(BeautifulSoup(html, "html.parser"), "https://www.gxrc.com/jobDetail/example")
    assert result is not None
    assert result["title"] == "业务经理"
    assert result["company"] == "示例科技有限公司"
    assert result["requirements"] == ["三年以上相关经验", "具备商务谈判能力。"]
    assert result["responsibilities"] == ["负责市场拓展", "推进项目落地。"]
    assert result["keywords"] == ["产教融合", "校企合作"]
    assert "推荐职位" not in result["raw_text"]


def test_mobile_gxrc_url_uses_desktop_detail_page(monkeypatch) -> None:
    captured = ""
    html = (
        '<html><body><h1>产品经理</h1><pre id="examineSensitiveWordsContent">'
        "任职要求：\n1.相关经验。\n岗位职责：\n1.产品规划。</pre></body></html>"
    )

    async def fake_fetch(url: str):
        nonlocal captured
        captured = url
        return httpx.Response(200, text=html, request=httpx.Request("GET", url))

    monkeypatch.setattr(services, "_fetch_public_page", fake_fetch)
    _, result = asyncio.run(
        services.fetch_job_page("https://m.gxrc.com/jobDetail/2c0e54d2-dd3b-4114-a1d2-58a54b3908dd")
    )

    assert captured == "https://www.gxrc.com/jobDetail/2c0e54d2-dd3b-4114-a1d2-58a54b3908dd"
    assert result is not None
    assert result["title"] == "产品经理"


def test_job_page_rejects_aliyun_waf_verification(monkeypatch) -> None:
    html = """
    <!doctype html><html><head><meta name="aliyun_waf_aa" content="token"></head>
    <body>为了更好的访问体验，请进行验证 appkey: "CF_APP_WAF"</body></html>
    """

    async def fake_fetch(url: str):
        return httpx.Response(200, text=html, request=httpx.Request("GET", url))

    monkeypatch.setattr(services, "_fetch_public_page", fake_fetch)
    with pytest.raises(ValueError, match="人机验证页"):
        asyncio.run(services.fetch_job_page("https://jobs.51job.com/nanning/171876687.html"))


def test_empty_jd_result_is_not_valid() -> None:
    with pytest.raises(ValueError, match=JD_EMPTY_ERROR):
        validate_jd_result(
            {
                "title": "",
                "company": "",
                "responsibilities": [],
                "requirements": [],
                "preferred": [],
                "keywords": [],
                "raw_text": "只有网页导航和广告，没有岗位正文",
            }
        )


def test_jd_result_rejects_verification_raw_text() -> None:
    with pytest.raises(ValueError, match="人机验证页"):
        validate_jd_result(
            {
                "title": "",
                "company": "",
                "responsibilities": [],
                "requirements": [],
                "preferred": [],
                "keywords": [],
                "raw_text": '为了更好的访问体验，请进行验证 appkey: "CF_APP_WAF"',
            }
        )


def test_invalid_model_json_is_retried_inside_chat_json() -> None:
    settings = Settings(storage_backend="memory", ai_mock=False, sensenova_max_retries=2)
    client = SenseNovaClient(settings, ObjectStore(settings))
    responses = iter(['{"title":', '{"title":"有效结果"}'])
    calls = 0

    async def fake_chat_once(*_args, **_kwargs):
        nonlocal calls
        calls += 1
        return next(responses)

    client._chat_once = fake_chat_once
    result = asyncio.run(client._chat_json("test-model", [{"role": "user", "content": "test"}]))
    assert result == {"title": "有效结果"}
    assert calls == 2


def test_persistent_invalid_model_json_has_stable_error() -> None:
    settings = Settings(storage_backend="memory", ai_mock=False, sensenova_max_retries=2)
    client = SenseNovaClient(settings, ObjectStore(settings))

    async def fake_chat_once(*_args, **_kwargs):
        return ""

    client._chat_once = fake_chat_once
    with pytest.raises(AIServiceError, match="AI 暂时没有返回有效内容") as captured:
        asyncio.run(client._chat_json("test-model", [{"role": "user", "content": "test"}]))
    assert captured.value.category == "empty_content"


def test_truncated_model_output_uses_compact_retry_and_records_metadata() -> None:
    settings = Settings(storage_backend="memory", ai_mock=False, sensenova_max_retries=2)
    client = SenseNovaClient(settings, ObjectStore(settings))
    calls = []

    async def fake_chat_once(_model, messages, _max_tokens):
        calls.append(messages)
        if len(calls) == 1:
            return "", {"model": "test-model", "finish_reason": "length"}
        return '{"title":"产品经理"}', {"model": "test-model", "finish_reason": "stop"}

    client._chat_once = fake_chat_once
    result = asyncio.run(client._chat_json("test-model", [{"role": "user", "content": "长岗位"}]))
    assert result == {"title": "产品经理"}
    assert len(calls) == 2
    assert "跳过分析过程" in calls[1][-1]["content"]


def test_public_page_connects_to_validated_ip_with_original_host_and_sni(monkeypatch) -> None:
    _mock_public_dns(monkeypatch)
    captured: list[httpx.Request] = []

    async def handler(request: httpx.Request) -> httpx.Response:
        captured.append(request)
        content = gzip.compress("<html>岗位</html>".encode())
        headers = {"content-type": "text/html; charset=utf-8", "content-encoding": "gzip"}
        return httpx.Response(200, headers=headers, stream=StaticStream(content))

    _mock_async_transport(monkeypatch, handler)
    original_url = "https://jobs.example.com/job?id=42"
    response = asyncio.run(services._fetch_public_page(original_url))

    assert str(captured[0].url) == f"https://{PUBLIC_TEST_IP}/job?id=42"
    assert captured[0].headers["host"] == "jobs.example.com"
    assert captured[0].extensions["sni_hostname"] == "jobs.example.com"
    assert str(response.url) == original_url
    assert response.text == "<html>岗位</html>"


def test_public_page_rejects_redirect_to_private_address(monkeypatch) -> None:
    _mock_public_dns(monkeypatch)
    requests: list[httpx.Request] = []

    async def handler(request: httpx.Request) -> httpx.Response:
        requests.append(request)
        return httpx.Response(302, headers={"location": "http://127.0.0.1/internal"})

    _mock_async_transport(monkeypatch, handler)
    with pytest.raises(ValueError, match="仅支持可公开访问"):
        asyncio.run(services._fetch_public_page("https://jobs.example.com/job"))
    assert len(requests) == 1


def test_public_page_stops_stream_after_two_megabytes(monkeypatch) -> None:
    _mock_public_dns(monkeypatch)
    stream = OversizedStream()

    async def handler(_request: httpx.Request) -> httpx.Response:
        return httpx.Response(200, headers={"content-type": "text/html"}, stream=stream)

    _mock_async_transport(monkeypatch, handler)
    with pytest.raises(ValueError, match="超过 2MB"):
        asyncio.run(services._fetch_public_page("https://jobs.example.com/job"))
    assert stream.closed is True


def test_public_page_timeout_has_stable_chinese_error(monkeypatch) -> None:
    _mock_public_dns(monkeypatch)

    async def handler(request: httpx.Request) -> httpx.Response:
        raise httpx.ReadTimeout("", request=request)

    _mock_async_transport(monkeypatch, handler)
    with pytest.raises(ValueError, match="读取岗位页面超时"):
        asyncio.run(services._fetch_public_page("https://jobs.example.com/job"))
